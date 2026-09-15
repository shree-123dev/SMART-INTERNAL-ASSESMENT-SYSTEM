import os
import re
from pptx import Presentation
from database import get_db_connection

def clean_text(text):
    if text is None:
        return ""
    return str(text).strip()

def is_usn_header(header):
    h = header.lower()
    return any(k in h for k in ["usn", "university seat", "seat no", "roll no", "reg no", "registration", "student id", "student_id"])

def is_name_header(header):
    h = header.lower()
    return any(k in h for k in ["name", "student name", "fullname", "student_name", "candidate"])

def is_ia_header(header, ia_type):
    h = header.lower()
    target = ia_type.lower() # 'ia1' or 'ia2'
    num = target.replace("ia", "") # '1' or '2'
    
    # Specific matching: "ia1", "ia 1", "ia-1", "internal 1", "cie 1", "test 1"
    if target in h or f"ia {num}" in h or f"ia-{num}" in h or f"internal {num}" in h or f"cie {num}" in h or f"test {num}" in h:
        return True
    return False

def is_generic_marks_header(header):
    h = header.lower()
    return any(k in h for k in ["marks", "score", "obtained", "total", "mark"])

def parse_pptx_ia_marks(pptx_file_or_path, subject_code, ia_type="IA1", max_marks=20.0):
    """
    Parses a PowerPoint (.pptx) presentation containing class IA marks.
    Extracts table or text data, matches USNs against students table,
    validates marks, detects updates vs new entries, and returns a structured preview.
    Accepts either a filesystem path (str) or a file-like object (BytesIO, FileStorage).
    """
    if isinstance(pptx_file_or_path, str):
        if not os.path.exists(pptx_file_or_path):
            return {"success": False, "error": "Presentation file not found."}

    try:
        prs = Presentation(pptx_file_or_path)
    except Exception as e:
        return {"success": False, "error": f"Failed to open presentation file: {str(e)}"}

    extracted_records = []
    seen_usns_in_file = set()

    # 1. SCAN TABLES IN ALL SLIDES
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.has_table:
                table = shape.table
                if len(table.rows) < 2:
                    continue

                # Analyze header row (row 0)
                headers = [clean_text(cell.text) for cell in table.rows[0].cells]
                
                usn_col_idx = None
                name_col_idx = None
                marks_col_idx = None

                # Find USN column
                for idx, h in enumerate(headers):
                    if is_usn_header(h):
                        usn_col_idx = idx
                        break

                # Find Name column
                for idx, h in enumerate(headers):
                    if idx != usn_col_idx and is_name_header(h):
                        name_col_idx = idx
                        break

                # Find specific IA marks column (e.g. IA1 or IA2)
                for idx, h in enumerate(headers):
                    if idx != usn_col_idx and idx != name_col_idx and is_ia_header(h, ia_type):
                        marks_col_idx = idx
                        break

                # Fallback to generic marks column if specific IA header not named
                if marks_col_idx is None:
                    for idx, h in enumerate(headers):
                        if idx != usn_col_idx and idx != name_col_idx and is_generic_marks_header(h):
                            marks_col_idx = idx
                            break

                # If we found at least USN and Marks column (or if table is 2-column [USN, Marks] or 3-column [USN, Name, Marks])
                if usn_col_idx is None and len(headers) >= 2:
                    # Check first row cell values: if first cell looks like USN (e.g. alphanumeric)
                    first_cell = headers[0].lower()
                    if "usn" in first_cell or any(c.isdigit() for c in headers[0]):
                        usn_col_idx = 0
                        if len(headers) == 2:
                            marks_col_idx = 1
                        elif len(headers) >= 3:
                            name_col_idx = 1
                            marks_col_idx = 2

                if usn_col_idx is not None and marks_col_idx is not None:
                    for row_idx in range(1, len(table.rows)):
                        row_cells = table.rows[row_idx].cells
                        if len(row_cells) <= max(usn_col_idx, marks_col_idx):
                            continue
                        
                        raw_usn = clean_text(row_cells[usn_col_idx].text).upper()
                        raw_name = clean_text(row_cells[name_col_idx].text) if name_col_idx is not None and len(row_cells) > name_col_idx else ""
                        raw_marks = clean_text(row_cells[marks_col_idx].text)

                        # Skip completely empty rows
                        if not raw_usn and not raw_marks:
                            continue

                        extracted_records.append({
                            "raw_usn": raw_usn,
                            "raw_name": raw_name,
                            "raw_marks": raw_marks,
                            "slide_no": slide_idx
                        })

    # 2. FALLBACK SCAN TEXT BOXES IF NO TABLES FOUND
    if not extracted_records:
        for slide_idx, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = clean_text(paragraph.text)
                        if not text:
                            continue
                        # Look for lines formatted like "23CS001, Student Name, 18" or "23CS001 | 18"
                        tokens = [clean_text(t) for t in re.split(r'[\t|,;]', text) if clean_text(t)]
                        if len(tokens) >= 2:
                            # Check if token 0 looks like USN
                            potential_usn = tokens[0].upper()
                            if re.match(r'^[0-9A-Z]{5,15}$', potential_usn):
                                raw_usn = potential_usn
                                if len(tokens) == 2:
                                    raw_name = ""
                                    raw_marks = tokens[1]
                                else:
                                    raw_name = tokens[1]
                                    raw_marks = tokens[2]
                                extracted_records.append({
                                    "raw_usn": raw_usn,
                                    "raw_name": raw_name,
                                    "raw_marks": raw_marks,
                                    "slide_no": slide_idx
                                })

    if not extracted_records:
        return {
            "success": False,
            "error": "Unable to recognize the marks table in the presentation. Please ensure your slide contains a table with 'USN' and 'Marks' or 'IA1'/'IA2' columns."
        }

    # 3. DATABASE VALIDATION & COMPARISON
    conn = get_db_connection()
    cursor = conn.cursor()

    parsed_rows = []
    summary = {
        "total_rows": len(extracted_records),
        "ready_to_import": 0,
        "to_update": 0,
        "unchanged": 0,
        "invalid": 0,
        "not_found": 0,
        "duplicates": 0
    }

    ia_field = "ia1" if ia_type.upper() == "IA1" else "ia2"

    for item in extracted_records:
        usn = item["raw_usn"]
        name_in_file = item["raw_name"]
        raw_marks = item["raw_marks"]
        slide_no = item["slide_no"]

        row_data = {
            "usn": usn,
            "name_in_file": name_in_file,
            "db_name": "",
            "department": "",
            "semester": "",
            "section": "",
            "raw_marks": raw_marks,
            "marks": None,
            "old_mark": None,
            "status": "valid_new",
            "status_label": "Ready to Import",
            "status_badge": "badge-success",
            "is_valid": False,
            "error_msg": "",
            "slide_no": slide_no
        }

        # Check empty USN
        if not usn:
            row_data["status"] = "invalid_mark"
            row_data["status_label"] = "Missing USN"
            row_data["status_badge"] = "badge-danger"
            row_data["error_msg"] = "USN is empty in presentation row."
            summary["invalid"] += 1
            parsed_rows.append(row_data)
            continue

        # Check duplicate USN in presentation
        if usn in seen_usns_in_file:
            row_data["status"] = "duplicate_in_file"
            row_data["status_label"] = "Duplicate USN in File"
            row_data["status_badge"] = "badge-danger"
            row_data["error_msg"] = f"USN '{usn}' appears multiple times in uploaded presentation."
            summary["duplicates"] += 1
            parsed_rows.append(row_data)
            continue
        
        seen_usns_in_file.add(usn)

        # Match USN in students table
        cursor.execute("""
            SELECT usn, fullname, department, semester, section
            FROM students
            WHERE usn = ?
        """, (usn,))
        student = cursor.fetchone()

        if not student:
            row_data["status"] = "student_not_found"
            row_data["status_label"] = "Student Not Found"
            row_data["status_badge"] = "badge-danger"
            row_data["error_msg"] = f"USN '{usn}' does not match any enrolled student record."
            summary["not_found"] += 1
            parsed_rows.append(row_data)
            continue

        row_data["db_name"] = student["fullname"]
        row_data["department"] = student["department"]
        row_data["semester"] = student["semester"]
        row_data["section"] = student["section"]

        # Validate numeric marks
        try:
            val = float(raw_marks)
        except (ValueError, TypeError):
            row_data["status"] = "invalid_mark"
            row_data["status_label"] = "Invalid (Non-numeric)"
            row_data["status_badge"] = "badge-danger"
            row_data["error_msg"] = f"Marks '{raw_marks}' is not a valid numeric value."
            summary["invalid"] += 1
            parsed_rows.append(row_data)
            continue

        # Validate bounds
        if val < 0:
            row_data["status"] = "invalid_mark"
            row_data["status_label"] = "Invalid (Negative)"
            row_data["status_badge"] = "badge-danger"
            row_data["error_msg"] = f"Marks ({val}) cannot be negative."
            summary["invalid"] += 1
            parsed_rows.append(row_data)
            continue

        if val > max_marks:
            row_data["status"] = "invalid_mark"
            row_data["status_label"] = f"Invalid (> {max_marks})"
            row_data["status_badge"] = "badge-danger"
            row_data["error_msg"] = f"Marks ({val}) exceeds maximum allowed ({max_marks})."
            summary["invalid"] += 1
            parsed_rows.append(row_data)
            continue

        row_data["marks"] = val

        # Query existing internal_marks
        cursor.execute(f"""
            SELECT id, {ia_field} as current_mark
            FROM internal_marks
            WHERE student_id = ? AND subject_id = ?
        """, (usn, subject_code))
        existing = cursor.fetchone()

        if existing and existing["current_mark"] is not None:
            curr_val = float(existing["current_mark"])
            row_data["old_mark"] = curr_val
            if curr_val == val:
                row_data["status"] = "unchanged"
                row_data["status_label"] = f"Unchanged ({curr_val})"
                row_data["status_badge"] = "badge-info"
                row_data["is_valid"] = True
                summary["unchanged"] += 1
            else:
                row_data["status"] = "valid_update"
                row_data["status_label"] = f"Update ({curr_val} → {val})"
                row_data["status_badge"] = "badge-warning"
                row_data["is_valid"] = True
                summary["to_update"] += 1
        else:
            row_data["status"] = "valid_new"
            row_data["status_label"] = "Ready to Import"
            row_data["status_badge"] = "badge-success"
            row_data["is_valid"] = True
            summary["ready_to_import"] += 1

        parsed_rows.append(row_data)

    conn.close()

    return {
        "success": True,
        "ia_type": ia_type.upper(),
        "subject_code": subject_code,
        "max_marks": max_marks,
        "summary": summary,
        "rows": parsed_rows
    }
