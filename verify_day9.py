import os
import sys
import sqlite3
import io
import unittest
from pptx import Presentation
from pptx.util import Inches, Pt
from werkzeug.security import generate_password_hash

# Ensure working directory in path
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))
from app import app
from database import get_db_connection, init_db
from ppt_importer import parse_pptx_ia_marks

def generate_test_pptx_table(data_rows, title="Internal Assessment 1 - Results"):
    """
    Generates a PPTX presentation with a table containing (USN, Student Name, Marks).
    data_rows: list of tuples (usn, name, mark_str)
    """
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    if slide.shapes.title:
        slide.shapes.title.text = title
        
    rows = len(data_rows) + 1
    cols = 3
    left = Inches(0.8)
    top = Inches(1.5)
    width = Inches(8.4)
    height = Inches(0.8 + 0.35 * len(data_rows))
    
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    
    # Headers
    table.cell(0, 0).text = "USN"
    table.cell(0, 1).text = "Student Name"
    table.cell(0, 2).text = "Marks (Max 20)"
    
    for i, (usn, name, mark) in enumerate(data_rows):
        table.cell(i + 1, 0).text = str(usn)
        table.cell(i + 1, 1).text = str(name)
        table.cell(i + 1, 2).text = str(mark)
        
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf

def generate_test_pptx_text_shapes(lines, title="IA Marks List"):
    """
    Generates a PPTX presentation where marks are in bullet/text shapes:
    e.g. "1MS24CS001, Alice Student, 18.5"
    """
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    if slide.shapes.title:
        slide.shapes.title.text = title
        
    txBox = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(8.0), Inches(5.0))
    tf = txBox.text_frame
    for i, line in enumerate(lines):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = line
        p.font.size = Pt(14)
        
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf

class TestDay9IAMarksSystem(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        init_db()
        conn = get_db_connection()
        cursor = conn.cursor()

        pw_hash = generate_password_hash("pass123")

        # Seed or Update Teachers
        # Teacher A (Assigned to CS401, CS402)
        cursor.execute("SELECT id FROM users WHERE email = 'teacherA@test.com'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO users (fullname, email, password, role) VALUES (?, ?, ?, 'teacher')",
                           ("Prof. Teacher Alpha", "teacherA@test.com", pw_hash))
        else:
            cursor.execute("UPDATE users SET password = ?, role = 'teacher' WHERE email = 'teacherA@test.com'", (pw_hash,))

        cursor.execute("SELECT id FROM teachers WHERE teacher_id = 'TCH_A'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO teachers (teacher_id, fullname, email, department) VALUES ('TCH_A', 'Prof. Teacher Alpha', 'teacherA@test.com', 'CSE')")

        # Teacher B (Assigned to CS403 only)
        cursor.execute("SELECT id FROM users WHERE email = 'teacherB@test.com'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO users (fullname, email, password, role) VALUES (?, ?, ?, 'teacher')",
                           ("Prof. Teacher Beta", "teacherB@test.com", pw_hash))
        else:
            cursor.execute("UPDATE users SET password = ?, role = 'teacher' WHERE email = 'teacherB@test.com'", (pw_hash,))

        cursor.execute("SELECT id FROM teachers WHERE teacher_id = 'TCH_B'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO teachers (teacher_id, fullname, email, department) VALUES ('TCH_B', 'Prof. Teacher Beta', 'teacherB@test.com', 'CSE')")

        # Admin
        cursor.execute("SELECT id FROM users WHERE email = 'admin@test.com'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO users (fullname, email, password, role) VALUES (?, ?, ?, 'admin')",
                           ("Administrator", "admin@test.com", generate_password_hash("admin123")))
        else:
            cursor.execute("UPDATE users SET password = ?, role = 'admin' WHERE email = 'admin@test.com'", (generate_password_hash("admin123"),))

        # Seed Students
        cursor.execute("SELECT id FROM users WHERE email = 'student1@test.com'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO users (fullname, email, password, role) VALUES (?, ?, ?, 'student')",
                           ("Alice Student", "student1@test.com", pw_hash))
        else:
            cursor.execute("UPDATE users SET password = ?, role = 'student' WHERE email = 'student1@test.com'", (pw_hash,))

        cursor.execute("SELECT id FROM students WHERE usn = '1MS24CS001'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO students (usn, fullname, email, department, semester, section) VALUES ('1MS24CS001', 'Alice Student', 'student1@test.com', 'CSE', 4, 'A')")

        cursor.execute("SELECT id FROM users WHERE email = 'student2@test.com'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO users (fullname, email, password, role) VALUES (?, ?, ?, 'student')",
                           ("Bob Student", "student2@test.com", pw_hash))
        else:
            cursor.execute("UPDATE users SET password = ?, role = 'student' WHERE email = 'student2@test.com'", (pw_hash,))

        cursor.execute("SELECT id FROM students WHERE usn = '1MS24CS002'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO students (usn, fullname, email, department, semester, section) VALUES ('1MS24CS002', 'Bob Student', 'student2@test.com', 'CSE', 4, 'B')")

        # Seed Subjects
        cursor.execute("SELECT id FROM subjects WHERE subject_code = 'CS401'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO subjects (subject_code, subject_name, department, semester) VALUES ('CS401', 'Data Structures Lab', 'CSE', 4)")

        cursor.execute("SELECT id FROM subjects WHERE subject_code = 'CS402'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO subjects (subject_code, subject_name, department, semester) VALUES ('CS402', 'Algorithms Lab', 'CSE', 4)")

        cursor.execute("SELECT id FROM subjects WHERE subject_code = 'CS403'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO subjects (subject_code, subject_name, department, semester) VALUES ('CS403', 'Database Systems', 'CSE', 4)")

        # Teacher Assignments:
        # Teacher A -> CS401, CS402
        # Teacher B -> CS403
        cursor.execute("SELECT id FROM teacher_subjects WHERE teacher_id = 'TCH_A' AND subject_id = 'CS401'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO teacher_subjects (teacher_id, subject_id) VALUES ('TCH_A', 'CS401')")

        cursor.execute("SELECT id FROM teacher_subjects WHERE teacher_id = 'TCH_A' AND subject_id = 'CS402'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO teacher_subjects (teacher_id, subject_id) VALUES ('TCH_A', 'CS402')")

        cursor.execute("SELECT id FROM teacher_subjects WHERE teacher_id = 'TCH_B' AND subject_id = 'CS403'")
        if not cursor.fetchone():
            cursor.execute("INSERT INTO teacher_subjects (teacher_id, subject_id) VALUES ('TCH_B', 'CS403')")

        # Seed practical session and Day 8 experiment marks for Alice to test continuous CIE calculation
        cursor.execute("DELETE FROM experiment_marks WHERE student_id = '1MS24CS001' AND subject_id = 'CS401'")
        cursor.execute("DELETE FROM practical_sessions WHERE teacher_id = 'TCH_A' AND subject_id = 'CS401'")
        cursor.execute("""
            INSERT INTO practical_sessions (teacher_id, subject_id, session_date, start_time, status, experiment_no, experiment_name)
            VALUES ('TCH_A', 'CS401', '2026-09-15', '09:00', 'active', 'Exp 1', 'Stack Implementation')
        """)
        sess_id = cursor.lastrowid

        cursor.execute("""
            INSERT INTO experiment_marks (student_id, teacher_id, subject_id, session_id, experiment_no, practical_marks, assignment_marks, recorded_date)
            VALUES ('1MS24CS001', 'TCH_A', 'CS401', ?, 'Exp 1', 9.0, 4.5, '2026-09-15'),
                   ('1MS24CS001', 'TCH_A', 'CS401', ?, 'Exp 2', 10.0, 5.0, '2026-09-15')
        """, (sess_id, sess_id))

        conn.commit()
        conn.close()

    def setUp(self):
        app.config['TESTING'] = True
        app.config['WTF_CSRF_ENABLED'] = False
        self.app = app.test_client()

    def login(self, email, password="pass123"):
        return self.app.post("/login", data={"email": email, "password": password}, follow_redirects=True)

    # 1. Database Schema Verification
    def test_01_database_schema(self):
        """Verify internal_marks table schema has all Day 9 upgraded columns."""
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("PRAGMA table_info(internal_marks)")
        cols = {row['name']: row['type'] for row in cur.fetchall()}
        conn.close()

        expected = ['id', 'student_id', 'subject_id', 'ia1', 'ia2', 'max_ia1', 'max_ia2', 'teacher_id', 'updated_at']
        for col in expected:
            self.assertIn(col, cols, f"Column {col} must exist in internal_marks table")
        print(" [PASS] test_01_database_schema: internal_marks schema upgraded and verified.")

    # 2. PPTX Parsing Engine Unit Tests
    def test_02_pptx_parsing_engine(self):
        """Verify PPTX parser handles valid rows, updates, unknown USNs, out-of-range marks, and duplicates."""
        test_rows = [
            ("1MS24CS001", "Alice Student", "18.5"),
            ("1MS24CS002", "Bob Student", "19.0"),
            ("1MS24CS002", "Bob Student", "17.0"),       # Duplicate row in same presentation
            ("1MS24CS999", "Unknown Ghost", "15.0"),     # Non-existent USN in DB
            ("1MS24CS001", "Alice Student", "25.0"),     # Out of range (>20) -> Marked invalid
        ]
        pptx_buf = generate_test_pptx_table(test_rows)
        
        parsed = parse_pptx_ia_marks(pptx_buf, "CS401", "IA1", 20.0)
        
        self.assertTrue(parsed['success'])
        self.assertEqual(parsed['summary']['ready_to_import'] + parsed['summary']['to_update'], 2)  # Alice & Bob
        self.assertEqual(parsed['summary']['not_found'], 1)        # 1MS24CS999
        self.assertEqual(parsed['summary']['duplicates'], 2)       # Duplicates in file
        
        # Test text-shapes format parsing (e.g. "1MS24CS001, Alice Student, 17.5")
        text_lines = [
            "1MS24CS001, Alice Student, 17.5",
            "1MS24CS002, Bob Student, 18.0"
        ]
        pptx_text_buf = generate_test_pptx_text_shapes(text_lines)
        parsed_text = parse_pptx_ia_marks(pptx_text_buf, "CS401", "IA2", 20.0)
        self.assertTrue(parsed_text['success'])
        self.assertEqual(parsed_text['summary']['ready_to_import'] + parsed_text['summary']['to_update'], 2)
        print(" [PASS] test_02_pptx_parsing_engine: Table and text-shape parsing, validation, and duplicate detection verified.")

    # 3. Method 1: Manual IA Marks Entry & Validation
    def test_03_manual_ia_entry_and_validation(self):
        """Verify teacher can manually enter and save IA marks, and invalid ranges are rejected."""
        self.login("teacherA@test.com")

        # Step A: Invalid mark (>20)
        res_invalid = self.app.post("/ia_marks/save_manual/CS401", data={
            "ia_type": "IA1",
            "mark_1MS24CS001": "24",
            "mark_1MS24CS002": "15"
        }, follow_redirects=True)
        self.assertIn(b"between 0 and 20", res_invalid.data)

        # Step B: Invalid negative mark
        res_neg = self.app.post("/ia_marks/save_manual/CS401", data={
            "ia_type": "IA1",
            "mark_1MS24CS001": "-2",
            "mark_1MS24CS002": "15"
        }, follow_redirects=True)
        self.assertIn(b"between 0 and 20", res_neg.data)

        # Step C: Valid entry for IA1
        res_valid_ia1 = self.app.post("/ia_marks/save_manual/CS401", data={
            "ia_type": "IA1",
            "mark_1MS24CS001": "18.0",
            "mark_1MS24CS002": "16.5"
        }, follow_redirects=True)
        self.assertEqual(res_valid_ia1.status_code, 200)
        self.assertIn(b"Successfully saved IA1 marks", res_valid_ia1.data)

        # Step D: Valid entry for IA2
        res_valid_ia2 = self.app.post("/ia_marks/save_manual/CS401", data={
            "ia_type": "IA2",
            "mark_1MS24CS001": "17.0",
            "mark_1MS24CS002": "19.0"
        }, follow_redirects=True)
        self.assertEqual(res_valid_ia2.status_code, 200)
        self.assertIn(b"Successfully saved IA2 marks", res_valid_ia2.data)

        # Verify in DB
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("SELECT ia1, ia2 FROM internal_marks WHERE student_id = '1MS24CS001' AND subject_id = 'CS401'")
        row = cur.fetchone()
        self.assertIsNotNone(row)
        self.assertEqual(row['ia1'], 18.0)
        self.assertEqual(row['ia2'], 17.0)
        conn.close()
        print(" [PASS] test_03_manual_ia_entry_and_validation: Manual entry saving and input validation bounds verified.")

    # 4. Method 2: Bulk PPT/PPTX Upload, Staging Preview & Confirm Commit
    def test_04_bulk_pptx_staging_and_confirm(self):
        """Verify bulk PPT upload displays interactive staging preview and confirms commit to database."""
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("DELETE FROM internal_marks WHERE subject_id = 'CS401'")
        cur.execute("""
            INSERT INTO internal_marks (student_id, subject_id, ia1, ia2, max_ia1, max_ia2, teacher_id)
            VALUES ('1MS24CS001', 'CS401', 18.0, 17.0, 20, 20, 'TCH_A')
        """)
        conn.commit()
        conn.close()

        self.login("teacherA@test.com")

        # Create PPTX updating Alice's IA1 to 19.5 and inserting Bob's IA1 as 18.5
        data = [
            ("1MS24CS001", "Alice Student", "19.5"),
            ("1MS24CS002", "Bob Student", "18.5"),
        ]
        pptx_buf = generate_test_pptx_table(data)

        # 1. Upload PPT
        res_upload = self.app.post("/ia_marks/upload_ppt/CS401", data={
            "ia_type": "IA1",
            "ppt_file": (pptx_buf, "CS401_IA1_Class.pptx")
        }, content_type="multipart/form-data", follow_redirects=True)
        self.assertEqual(res_upload.status_code, 200)
        self.assertIn(b"PPT Import Preview", res_upload.data)
        self.assertIn(b"1MS24CS001", res_upload.data)
        self.assertIn(b"19.5", res_upload.data)

        # 2. Confirm Import
        res_confirm = self.app.post("/ia_marks/confirm_import", follow_redirects=True)
        self.assertEqual(res_confirm.status_code, 200)
        self.assertIn(b"Bulk Import Complete", res_confirm.data)

        # 3. Verify DB has the new mark
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("""
            SELECT im.ia1, im.ia2 
            FROM internal_marks im 
            WHERE im.subject_id = 'CS401' AND im.student_id = '1MS24CS001'
        """)
        row = cur.fetchone()
        self.assertIsNotNone(row)
        self.assertEqual(row['ia1'], 19.5)
        # IA2 should be preserved (17.0)
        self.assertEqual(row['ia2'], 17.0)
        conn.close()
        print(" [PASS] test_04_bulk_pptx_staging_and_confirm: PPT staging preview, update detection, and commit verified.")

    # 5. Bulk PPT Import Cancel Action
    def test_05_bulk_pptx_cancel_staging(self):
        """Verify teacher can cancel a staged PPT import without modifying database."""
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("DELETE FROM internal_marks WHERE subject_id = 'CS401' AND student_id = '1MS24CS001'")
        cur.execute("""
            INSERT INTO internal_marks (student_id, subject_id, ia1, ia2, max_ia1, max_ia2, teacher_id)
            VALUES ('1MS24CS001', 'CS401', 19.5, 17.0, 20, 20, 'TCH_A')
        """)
        conn.commit()
        conn.close()

        self.login("teacherA@test.com")

        # Create PPTX with 10.0
        data = [("1MS24CS001", "Alice Student", "10.0")]
        pptx_buf = generate_test_pptx_table(data)

        # Upload
        self.app.post("/ia_marks/upload_ppt/CS401", data={
            "ia_type": "IA1",
            "ppt_file": (pptx_buf, "CancelTest.pptx")
        }, content_type="multipart/form-data", follow_redirects=True)

        # Cancel
        res_cancel = self.app.get("/ia_marks/cancel_import/CS401", follow_redirects=True)
        self.assertEqual(res_cancel.status_code, 200)
        self.assertIn(b"cancelled", res_cancel.data.lower())

        # Verify DB still has 19.5 (not overwritten by 10.0)
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("""
            SELECT im.ia1 FROM internal_marks im 
            WHERE im.subject_id = 'CS401' AND im.student_id = '1MS24CS001'
        """)
        row = cur.fetchone()
        self.assertEqual(row['ia1'], 19.5)
        conn.close()
        print(" [PASS] test_05_bulk_pptx_cancel_staging: Staging cancellation leaves database unaltered.")

    # 6. CIE Dynamic Calculations Engine
    def test_06_cie_dynamic_calculations(self):
        """
        Verify Continuous Internal Evaluation (CIE) calculation:
        - IA1 = 19.5, IA2 = 17.0 -> IA Avg = (19.5 + 17.0)/2 = 18.25 (/20)
        - Practical marks = 9.0, 10.0 -> Lab Avg = 9.5 (/10)
        - Assignment marks = 4.5, 5.0 -> Assign Avg = 4.75 (/5)
        - Total CIE = 18.25 + 9.5 + 4.75 = 32.5 (/35)
        """
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("DELETE FROM internal_marks WHERE subject_id = 'CS401' AND student_id = '1MS24CS001'")
        cur.execute("""
            INSERT INTO internal_marks (student_id, subject_id, ia1, ia2, max_ia1, max_ia2, teacher_id)
            VALUES ('1MS24CS001', 'CS401', 19.5, 17.0, 20, 20, 'TCH_A')
        """)
        conn.commit()
        conn.close()

        self.login("teacherA@test.com")
        res = self.app.get("/ia_marks/CS401")
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"18.25", res.data)  # IA Avg
        self.assertIn(b"9.5", res.data)    # Lab Avg
        self.assertIn(b"4.75", res.data)   # Assign Avg
        self.assertIn(b"32.5", res.data)   # Total CIE (/35)
        print(" [PASS] test_06_cie_dynamic_calculations: CIE dynamic score computation accurately evaluated across all components.")

    # 7. Student View and Read-Only Protection
    def test_07_student_view_and_protection(self):
        """Verify student can view their calculated CIE breakdown and cannot access teacher IA routes."""
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("DELETE FROM internal_marks WHERE subject_id = 'CS401' AND student_id = '1MS24CS001'")
        cur.execute("""
            INSERT INTO internal_marks (student_id, subject_id, ia1, ia2, max_ia1, max_ia2, teacher_id)
            VALUES ('1MS24CS001', 'CS401', 19.5, 17.0, 20, 20, 'TCH_A')
        """)
        conn.commit()
        conn.close()

        self.login("student1@test.com")
        res = self.app.get("/student")
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Continuous Internal Assessment (CIE)", res.data)
        self.assertIn(b"CS401", res.data)
        self.assertIn(b"32.5", res.data) # Total CIE

        # Student cannot access teacher IA Gradebook routes (redirected to student dashboard)
        res_forbidden = self.app.get("/ia_marks/CS401", follow_redirects=False)
        self.assertEqual(res_forbidden.status_code, 302)

        # Student cannot submit manual marks
        res_post_forbidden = self.app.post("/ia_marks/save_manual/CS401", data={"ia_type": "IA1", "mark_1MS24CS001": "20"}, follow_redirects=False)
        self.assertEqual(res_post_forbidden.status_code, 302)
        print(" [PASS] test_07_student_view_and_protection: Student portal displays calculated CIE, blocks teacher endpoints.")

    # 8. Teacher Subject Isolation (Teacher A cannot touch Teacher B's subjects)
    def test_08_teacher_subject_isolation(self):
        """Verify Teacher A cannot access or import marks for Teacher B's assigned subject (CS403)."""
        self.login("teacherA@test.com")
        
        # Try to view CS403 Gradebook
        res = self.app.get("/ia_marks/CS403", follow_redirects=True)
        self.assertIn(b"Unauthorized", res.data)

        # Try to upload PPT for CS403
        data = [("1MS24CS001", "Alice Student", "20.0")]
        pptx_buf = generate_test_pptx_table(data)
        res_upload = self.app.post("/ia_marks/upload_ppt/CS403", data={
            "ia_type": "IA1",
            "ppt_file": (pptx_buf, "Hack.pptx")
        }, content_type="multipart/form-data", follow_redirects=True)
        self.assertIn(b"Unauthorized", res_upload.data)
        print(" [PASS] test_08_teacher_subject_isolation: Cross-teacher subject access strictly blocked.")

    # 9. Admin Marks Institutional Overview
    def test_09_admin_marks_roster(self):
        """Verify Admin can view all student IA marks and CIE totals across subjects."""
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("DELETE FROM internal_marks WHERE subject_id = 'CS401' AND student_id = '1MS24CS001'")
        cur.execute("""
            INSERT INTO internal_marks (student_id, subject_id, ia1, ia2, max_ia1, max_ia2, teacher_id)
            VALUES ('1MS24CS001', 'CS401', 19.5, 17.0, 20, 20, 'TCH_A')
        """)
        conn.commit()
        conn.close()

        self.login("admin@test.com", "admin123")
        res = self.app.get("/admin/marks")
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Academic Marks", res.data)
        self.assertIn(b"1MS24CS001", res.data)
        self.assertIn(b"Alice Student", res.data)
        self.assertIn(b"32.5", res.data)
        print(" [PASS] test_09_admin_marks_roster: Admin institution-wide master marks roster fully operational.")

    # 10. Regression and Day 1-8 Preservation
    def test_10_day1_to_8_preservation(self):
        """Verify Day 1-8 tables, practical sessions, attendance, QR codes, and experiment marks remain intact."""
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("SELECT name FROM sqlite_master WHERE type='table'")
        tables = [r['name'] for r in cur.fetchall()]
        
        required_tables = [
            'users', 'students', 'teachers', 'subjects', 'teacher_subjects',
            'attendance', 'practical_sessions', 'experiment_marks', 'internal_marks'
        ]
        for tbl in required_tables:
            self.assertIn(tbl, tables, f"Table {tbl} must be present in database")
            
        # Verify Day 8 experiment marks intact
        cur.execute("SELECT COUNT(*) as count FROM experiment_marks")
        count = cur.fetchone()['count']
        self.assertGreaterEqual(count, 1)
        conn.close()
        print(" [PASS] test_10_day1_to_8_preservation: All Day 1-8 entities and relational integrity fully preserved.")

if __name__ == '__main__':
    unittest.main(verbosity=2)
